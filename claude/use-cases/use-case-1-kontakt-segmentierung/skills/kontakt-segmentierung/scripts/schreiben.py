#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
schreiben.py: zugeordnete Kontakte als Excel-Datei und Markdown-Zusammenfassung schreiben.

Aufruf:
  python3 schreiben.py --arbeit <ordner> [--segmente <segmente.md>] --ausgabe <ordner> [--datum JJJJ-MM-TT] [--vorschau] [--cloud-ok]

Ohne --segmente gilt kontakte/segmente.md im Arbeitsordner, sonst die Fassung im Skill-Ordner.

Erwartet im Arbeitsordner:
  kontakte.json          aus einlesen.py
  zuordnung-*.json       je Block eine Liste [{"id": 1, "segment": "...", "prioritaet": "A", "begruendung": "..."}]
  bausteine.json         {"anrede": "Sie", "bausteine": {"<Segment>": "<Text>", ...}}
                         (optional zusätzlich der Schlüssel "Bisherige Partner" für Firmen aus der Partnerliste)
  bestand.json           optional, aus bestand.py (bisherige Partner, Aussteller, Sponsoren)

Mit --vorschau wird nur geprüft und eine Übersicht als Text ausgegeben, es
entsteht keine Datei. Ohne --vorschau entstehen
  <ausgabe>/segmentierung-<datum>.xlsx  (Blätter Kontakte, Firmen, Textbausteine)
  <ausgabe>/segmentierung-<datum>.md    (Zusammenfassung mit den Top-20-Firmen)
  <ausgabe>/segmentierung-<datum>.html  (Übersicht zum Anschauen im Browser, markiert A-Kontakte und bisherige Partner)
Mit --reiter "<Titel>" landet die HTML-Übersicht zusätzlich als eigener Reiter im AIOS-Dashboard
(~/.claude/dashboard/public/reiter/<slug>.html plus <slug>.json). Ein späterer Lauf mit demselben Titel
ersetzt den Reiter, das Dashboard zeigt den neuen Stand beim nächsten Datenabruf.
Bestehende Dateien werden nie überschrieben, dann entsteht -2, -3 usw.
Die Excel-Datei enthält keinen Autor, Bearbeiter oder Anwendungsnamen.
"""

import argparse
import datetime as dt
import glob
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from einlesen import cloud_hinweis, firma_schluessel, regeldatei  # noqa: E402

PRIOS = ("A", "B", "C")
BESTAND = "Bisherige Partner"  # Schlüssel für den Textbaustein an Firmen aus der Partnerliste
# Angaben im Feld Firma, die keine Firma sind, landen in der Firmenübersicht unter (ohne Firma)
KEINE_FIRMA = {"selbststandig", "selbstständig", "freiberuflich", "freelance", "freelancer", "self employed",
               "selbstandig", "freiberufler", "freiberuflerin", "n a", "keine", "none", "ohne"}


def fehler(text, details=None, code=2):
    aus = {"fehler": text}
    if details:
        aus["details"] = details
    print(json.dumps(aus, ensure_ascii=False, indent=1))
    sys.exit(code)


def zahl(n):
    return "{:,}".format(n).replace(",", ".")


def datum_de(iso):
    try:
        return dt.date.fromisoformat(iso).strftime("%d.%m.%Y")
    except (TypeError, ValueError):
        return iso or ""


def segmente_lesen(pfad):
    namen = []
    for zeile in open(pfad, encoding="utf-8"):
        m = re.match(r"^##\s+Segment:\s*(.+?)\s*$", zeile)
        if m:
            namen.append(m.group(1))
    if not namen:
        fehler("In %s steht keine Überschrift der Form '## Segment: <Name>'." % pfad)
    if "Sonstiges" not in namen:
        namen.append("Sonstiges")
    return namen


def segment_finden(wert, segmente):
    w = str(wert or "").strip().lower()
    for s in segmente:
        if s.lower() == w:
            return s
    treffer = [s for s in segmente if s.lower().startswith(w) or w.startswith(s.lower().split(" und ")[0])]
    return treffer[0] if len(treffer) == 1 and w else None


def laden(args):
    kpfad = os.path.join(args.arbeit, "kontakte.json")
    if not os.path.isfile(kpfad):
        fehler("kontakte.json fehlt im Arbeitsordner. Zuerst einlesen.py ausführen.")
    daten = json.load(open(kpfad, encoding="utf-8"))
    kontakte = {k["id"]: k for k in daten["kontakte"]}
    segmente = segmente_lesen(args.segmente)

    probleme, zuordnung = [], {}
    dateien = sorted(glob.glob(os.path.join(args.arbeit, "zuordnung-*.json")))
    if not dateien:
        fehler("Keine zuordnung-*.json im Arbeitsordner gefunden.")
    for datei in dateien:
        try:
            inhalt = json.load(open(datei, encoding="utf-8"))
        except ValueError as e:
            fehler("%s ist kein gültiges JSON: %s" % (os.path.basename(datei), e))
        if isinstance(inhalt, dict):
            inhalt = [dict(v, id=int(k)) for k, v in inhalt.items()]
        for z in inhalt:
            try:
                kid = int(z.get("id"))
            except (TypeError, ValueError):
                probleme.append("%s: Eintrag ohne gültige id" % os.path.basename(datei))
                continue
            if kid not in kontakte:
                probleme.append("id %d gibt es nicht" % kid)
                continue
            if kid in zuordnung:
                probleme.append("id %d doppelt zugeordnet" % kid)
            seg = segment_finden(z.get("segment"), segmente)
            prio = str(z.get("prioritaet", "")).strip().upper()
            if not seg:
                probleme.append("id %d: Segment '%s' unbekannt" % (kid, z.get("segment")))
            if prio not in PRIOS:
                probleme.append("id %d: Priorität '%s' ist nicht A, B oder C" % (kid, z.get("prioritaet")))
            zuordnung[kid] = {"segment": seg, "prioritaet": prio,
                              "begruendung": re.sub(r"\s+", " ", str(z.get("begruendung", "")).strip())}
    fehlend = [i for i in kontakte if i not in zuordnung]
    if fehlend:
        probleme.append("%d Kontakte ohne Zuordnung, etwa id %s" % (len(fehlend), ", ".join(map(str, fehlend[:15]))))
    if probleme:
        fehler("Zuordnung unvollständig oder fehlerhaft.", probleme)

    bpfad = os.path.join(args.arbeit, "bausteine.json")
    bausteine = json.load(open(bpfad, encoding="utf-8")) if os.path.isfile(bpfad) else {"anrede": "", "bausteine": {}}
    texte = {}
    for name, text in (bausteine.get("bausteine") or {}).items():
        if str(name).strip().lower() in ("bisherige partner", "bestand", "reaktivierung"):
            texte[BESTAND] = str(text).strip()
            continue
        seg = segment_finden(name, segmente)
        if not seg:
            fehler("bausteine.json: Segment '%s' unbekannt." % name)
        texte[seg] = str(text).strip()

    bestand = bestand_laden(args.arbeit)
    zeilen = []
    for kid, k in kontakte.items():
        z = dict(k, **zuordnung[kid])
        p = bestand["partner"].get(firma_schluessel(k["firma"]))
        z["bestand"] = bestand_kurz(p) if p else ""
        zeilen.append(z)
    rang = {s: i for i, s in enumerate(segmente)}
    zeilen.sort(key=lambda z: (PRIOS.index(z["prioritaet"]), rang[z["segment"]],
                               z["firma"].lower() or "~", z["nachname"].lower()))
    meta = dict(daten.get("meta", {}))
    meta["bestand"] = bestand
    return meta, zeilen, segmente, bausteine.get("anrede", ""), texte


def bestand_laden(arbeit):
    """bestand.json aus bestand.py, wenn vorhanden. Gibt {"quelle": ..., "partner": {schluessel: eintrag}} zurück."""
    pfad = os.path.join(arbeit, "bestand.json")
    if not os.path.isfile(pfad):
        return {"quelle": "", "partner": {}}
    try:
        daten = json.load(open(pfad, encoding="utf-8"))
    except ValueError as e:
        fehler("bestand.json ist kein gültiges JSON: %s" % e)
    partner = {}
    for p in daten.get("partner") or []:
        schluessel = p.get("schluessel") or firma_schluessel(p.get("firma"))
        if schluessel:
            partner[schluessel] = p
    return {"quelle": daten.get("quelle", ""), "partner": partner}


def bestand_kurz(p):
    """Kurzform für Tabellen, etwa 'Aussteller 2025' oder 'Partner'."""
    if not p:
        return ""
    teile = [t for t in (p.get("art"), p.get("jahr")) if t]
    return " ".join(teile) if teile else "ja"


def firmen_bilden(zeilen, segmente):
    rang = {s: i for i, s in enumerate(segmente)}
    gruppen = {}
    for z in zeilen:
        schluessel = firma_schluessel(z["firma"]) or "~ohne"
        if schluessel in KEINE_FIRMA:
            schluessel = "~ohne"
        gruppen.setdefault(schluessel, []).append(z)
    firmen = []
    for schluessel, leute in gruppen.items():
        bester = min(leute, key=lambda z: (PRIOS.index(z["prioritaet"]), rang[z["segment"]]))
        firmen.append({
            "firma": leute[0]["firma"] if schluessel != "~ohne" else "(ohne Firma)",
            "kontakte": len(leute),
            "a": sum(1 for z in leute if z["prioritaet"] == "A"),
            "b": sum(1 for z in leute if z["prioritaet"] == "B"),
            "segment": bester["segment"],
            "prioritaet": bester["prioritaet"],
            "kontakt": ("%s %s" % (bester["vorname"], bester["nachname"])).strip(),
            "position": bester["position"],
            "ohne": schluessel == "~ohne",
            "bestand": leute[0].get("bestand", ""),
        })
    firmen.sort(key=lambda f: (f["ohne"], PRIOS.index(f["prioritaet"]), rang[f["segment"]],
                               -f["a"], -f["kontakte"], f["firma"].lower()))
    return firmen


def verteilung(zeilen, segmente):
    tab = {s: {p: 0 for p in PRIOS} for s in segmente}
    for z in zeilen:
        tab[z["segment"]][z["prioritaet"]] += 1
    return tab


def vorschau(meta, zeilen, segmente, anrede, texte):
    tab = verteilung(zeilen, segmente)
    print("Quelle: %s, %s Kontakte" % (meta.get("quelle", "?"), zahl(len(zeilen))))
    print("")
    print("Segment | A | B | C | Gesamt")
    for s in segmente:
        t = tab[s]
        print("%s | %d | %d | %d | %d" % (s, t["A"], t["B"], t["C"], sum(t.values())))
    print("")
    print("Top-20-Firmen: Rang | Firma | Kontakte | davon A | bestes Segment | beste Priorität | wichtigster Kontakt")
    for i, f in enumerate([f for f in firmen_bilden(zeilen, segmente) if not f["ohne"]][:20], 1):
        print("%d | %s | %d | %d | %s | %s | %s (%s)" % (i, f["firma"], f["kontakte"], f["a"], f["segment"],
                                                       f["prioritaet"], f["kontakt"], f["position"]))
    print("")
    a = [z for z in zeilen if z["prioritaet"] == "A"]
    print("Priorität A, %d Kontakte (bis zu 30 gezeigt): Name | Position | Firma | Segment | Begründung" % len(a))
    for z in a[:30]:
        print("%s %s | %s | %s | %s | %s" % (z["vorname"], z["nachname"], z["position"], z["firma"],
                                            z["segment"], z["begruendung"]))
    print("")
    fehlt = [s for s in segmente if sum(tab[s].values()) and s not in texte]
    print("Textbausteine: %d (Anrede %s)%s" % (len(texte), anrede or "offen",
                                               ", fehlen für " + ", ".join(fehlt) if fehlt else ""))
    bestand = meta.get("bestand") or {}
    if bestand.get("partner"):
        firmen = firmen_bilden(zeilen, segmente)
        treffer = [f for f in firmen if f["bestand"]]
        print("")
        print("Bisherige Partner (%s): %d von %d Firmen auch im Export: Firma | Bestand | Kontakte | davon A"
              % (bestand.get("quelle", "?"), len(treffer), len(bestand["partner"])))
        for f in treffer:
            print("%s | %s | %d | %d" % (f["firma"], f["bestand"], f["kontakte"], f["a"]))
        im_export = {firma_schluessel(z["firma"]) for z in zeilen}
        ohne = [p for s, p in bestand["partner"].items() if s not in im_export]
        if ohne:
            print("Partner ohne Kontakt im Export: " + ", ".join(p["firma"] for p in ohne))
        if BESTAND not in texte:
            print("Textbaustein für bisherige Partner fehlt noch (Schlüssel 'Bisherige Partner' in bausteine.json).")


def freier_name(ordner, basis):
    n = 1
    while True:
        endung = "" if n == 1 else "-%d" % n
        if not (os.path.exists(os.path.join(ordner, basis + endung + ".xlsx")) or
                os.path.exists(os.path.join(ordner, basis + endung + ".md")) or
                os.path.exists(os.path.join(ordner, basis + endung + ".html"))):
            return basis + endung
        n += 1


def metadaten_leeren(pfad):
    """Entfernt Autor, Bearbeiter und Anwendungsnamen aus der fertigen xlsx-Datei."""
    app_xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
               '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
               'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
               '<Application></Application></Properties>')
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(pfad))
    os.close(fd)
    with zipfile.ZipFile(pfad) as alt, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as neu:
        for eintrag in alt.infolist():
            daten = alt.read(eintrag.filename)
            if eintrag.filename == "docProps/app.xml":
                daten = app_xml.encode("utf-8")
            elif eintrag.filename == "docProps/core.xml":
                text = daten.decode("utf-8")
                text = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator></dc:creator>", text, flags=re.S)
                text = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy></cp:lastModifiedBy>",
                              text, flags=re.S)
                daten = text.encode("utf-8")
            neu.writestr(eintrag, daten)
    shutil.move(tmp, pfad)


def excel_schreiben(pfad, meta, zeilen, segmente, anrede, texte, firmen):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        fehler("Für Excel-Dateien fehlt das Zusatzpaket openpyxl. Installieren mit: python3 -m pip install openpyxl "
              "(unter Windows py -3 -m pip install openpyxl).")

    fett = Font(bold=True, color="FFFFFF")
    kopf_fuellung = PatternFill("solid", fgColor="1F3A5F")
    prio_fuellung = {"A": PatternFill("solid", fgColor="C6EFCE"),
                     "B": PatternFill("solid", fgColor="FFEB9C"),
                     "C": PatternFill("solid", fgColor="EDEDED")}
    umbruch = Alignment(wrap_text=True, vertical="top")
    oben = Alignment(vertical="top")

    def text(zelle, wert):
        zelle.value = wert
        if isinstance(wert, str):
            zelle.data_type = "s"  # nie als Formel auswerten, auch wenn der Text mit = beginnt

    def blatt(ws, kopf, breiten):
        for j, (titel, breite) in enumerate(zip(kopf, breiten), 1):
            c = ws.cell(row=1, column=j)
            text(c, titel)
            c.font, c.fill, c.alignment = fett, kopf_fuellung, Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = breite
        ws.freeze_panes = "A2"

    wb = Workbook()
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.properties.title = None

    # Blatt 1: Kontakte
    ws = wb.active
    ws.title = "Kontakte"
    mit_bestand = bool((meta.get("bestand") or {}).get("partner"))
    kopf = ["Priorität", "Segment", "Vorname", "Nachname", "Firma", "Bisheriger Partner", "Position", "Begründung",
            "Verbunden seit", "Profil-Link", "E-Mail (aus Export)", "Status (für dich)"]
    breiten = [10, 30, 16, 18, 30, 18, 40, 50, 14, 38, 30, 18]
    if not mit_bestand:
        kopf.pop(5)
        breiten.pop(5)
    blatt(ws, kopf, breiten)
    bestand_fuellung = PatternFill("solid", fgColor="DDEBF7")
    for i, z in enumerate(zeilen, 2):
        werte = [z["prioritaet"], z["segment"], z["vorname"], z["nachname"], z["firma"]]
        if mit_bestand:
            werte.append(z.get("bestand", ""))
        werte += [z["position"], z["begruendung"]]
        for j, w in enumerate(werte, 1):
            text(ws.cell(row=i, column=j), w)
            ws.cell(row=i, column=j).alignment = oben
        ws.cell(row=i, column=1).fill = prio_fuellung[z["prioritaet"]]
        ws.cell(row=i, column=1).alignment = Alignment(horizontal="center", vertical="top")
        if mit_bestand and z.get("bestand"):
            ws.cell(row=i, column=6).fill = bestand_fuellung
        n = len(werte)
        c = ws.cell(row=i, column=n + 1)
        try:
            c.value = dt.date.fromisoformat(z["verbunden_seit"])
            c.number_format = "DD.MM.YYYY"
        except (TypeError, ValueError):
            text(c, z["verbunden_seit"])
        c = ws.cell(row=i, column=n + 2)
        text(c, z["url"])
        if re.match(r"^https?://", z["url"] or ""):
            c.hyperlink = z["url"]
            c.font = Font(color="0563C1", underline="single")
        text(ws.cell(row=i, column=n + 3), z["email"])
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(kopf)), max(len(zeilen) + 1, 2))

    # Blatt 2: Firmen
    ws = wb.create_sheet("Firmen")
    kopf = ["Rang", "Firma", "Bisheriger Partner", "Kontakte", "davon A", "davon B", "Bestes Segment",
            "Beste Priorität", "Wichtigster Kontakt", "Position"]
    breiten = [7, 34, 18, 11, 10, 10, 32, 12, 26, 40]
    if not mit_bestand:
        kopf.pop(2)
        breiten.pop(2)
    blatt(ws, kopf, breiten)
    for i, f in enumerate(firmen, 2):
        werte = [i - 1 if not f["ohne"] else None, f["firma"]]
        if mit_bestand:
            werte.append(f.get("bestand", ""))
        werte += [f["kontakte"], f["a"], f["b"], f["segment"], f["prioritaet"], f["kontakt"], f["position"]]
        for j, w in enumerate(werte, 1):
            c = ws.cell(row=i, column=j)
            text(c, w)
            if isinstance(w, int):
                c.number_format = "#,##0"
        if mit_bestand and f.get("bestand"):
            ws.cell(row=i, column=3).fill = bestand_fuellung
        sp = len(werte) - 2
        ws.cell(row=i, column=sp).fill = prio_fuellung[f["prioritaet"]]
        ws.cell(row=i, column=sp).alignment = Alignment(horizontal="center")
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(kopf)), max(len(firmen) + 1, 2))

    # Blatt 3: Textbausteine
    ws = wb.create_sheet("Textbausteine")
    kopf = ["Segment", "Anrede", "Kontakte im Segment", "Textbaustein (LinkedIn-Nachricht)", "Zeichen"]
    blatt(ws, kopf, [30, 9, 12, 90, 9])
    tab = verteilung(zeilen, segmente)
    zeile = 2
    for s in list(segmente) + [BESTAND]:
        if s not in texte:
            continue
        t = texte[s]
        anzahl = sum(1 for z in zeilen if z.get("bestand")) if s == BESTAND else sum(tab[s].values())
        for j, w in enumerate([s, anrede, anzahl, t, len(t)], 1):
            c = ws.cell(row=zeile, column=j)
            text(c, w)
            c.alignment = umbruch
            if isinstance(w, int):
                c.number_format = "#,##0"
        ws.row_dimensions[zeile].height = min(409, 15 * (t.count("\n") + len(t) // 85 + 2))
        zeile += 1
    hinweis = ws.cell(row=zeile + 1, column=1)
    text(hinweis, "Platzhalter in eckigen Klammern vor dem Senden ersetzen. Jede Nachricht schickst du selbst, "
                  "nichts wird automatisch versendet.")
    hinweis.font = Font(italic=True, color="666666")

    wb.save(pfad)
    metadaten_leeren(pfad)


def markdown_schreiben(pfad, meta, zeilen, segmente, firmen, datum, xlsx_name):
    tab = verteilung(zeilen, segmente)
    z = []
    z.append("# Kontakt-Segmentierung für den ADAN Career Day, %s" % datum_de(datum))
    z.append("")
    ausgeschlossen = sum((meta.get("ausgeschlossen") or {}).values())
    z.append("Quelle %s, eingelesen am %s. %s Kontakte aus %s Firmen segmentiert%s. Die vollständige Liste steht in `%s`."
             % (meta.get("quelle", "?"), datum_de(meta.get("eingelesen_am", "")), zahl(len(zeilen)),
                zahl(len([f for f in firmen if not f["ohne"]])),
                (", %s Kontakt über die Ausschlussliste weggelassen" if ausgeschlossen == 1 else
                 ", %s Kontakte über die Ausschlussliste weggelassen") % zahl(ausgeschlossen) if ausgeschlossen else "",
                xlsx_name))
    z.append("")
    z.append("## Verteilung")
    z.append("")
    z.append("| Segment | A | B | C | Gesamt |")
    z.append("|---|---:|---:|---:|---:|")
    for s in segmente:
        t = tab[s]
        z.append("| %s | %s | %s | %s | %s |" % (s, zahl(t["A"]), zahl(t["B"]), zahl(t["C"]), zahl(sum(t.values()))))
    summe = {p: sum(tab[s][p] for s in segmente) for p in PRIOS}
    z.append("| **Gesamt** | **%s** | **%s** | **%s** | **%s** |" % (zahl(summe["A"]), zahl(summe["B"]),
                                                                     zahl(summe["C"]), zahl(len(zeilen))))
    z.append("")
    z.append("## Top-20-Firmen")
    z.append("")
    z.append("Sortiert nach bester Priorität, dann Segment, dann Zahl der A-Kontakte und Kontakte insgesamt.")
    z.append("")
    z.append("| Rang | Firma | Kontakte | davon A | Bestes Segment | Wichtigster Kontakt (Position) |")
    z.append("|---:|---|---:|---:|---|---|")
    for i, f in enumerate([f for f in firmen if not f["ohne"]][:20], 1):
        z.append("| %d | %s | %s | %s | %s (%s) | %s |" % (i, f["firma"].replace("|", "/"), zahl(f["kontakte"]),
                                                          zahl(f["a"]), f["segment"], f["prioritaet"],
                                                          (f["position"] or "ohne Position").replace("|", "/")))
    z.append("")
    bestand = meta.get("bestand") or {}
    if bestand.get("partner"):
        treffer = [f for f in firmen if f["bestand"]]
        im_export = {firma_schluessel(x["firma"]) for x in zeilen}
        ohne = [p for k, p in bestand["partner"].items() if k not in im_export]
        z.append("## Bisherige Partner (Reaktivierung)")
        z.append("")
        z.append("Partnerliste %s mit %s Firmen. %s davon haben Kontakte im Export, %s nicht."
                 % (bestand.get("quelle", "?"), zahl(len(bestand["partner"])), zahl(len(treffer)), zahl(len(ohne))))
        z.append("")
        if treffer:
            z.append("| Firma | Bisher | Kontakte | davon A | Wichtigster Kontakt (Position) |")
            z.append("|---|---|---:|---:|---|")
            for f in treffer:
                z.append("| %s | %s | %s | %s | %s (%s) |" % (f["firma"].replace("|", "/"), f["bestand"], zahl(f["kontakte"]),
                                                          zahl(f["a"]), f["kontakt"], (f["position"] or "ohne Position").replace("|", "/")))
            z.append("")
        if ohne:
            z.append("Ohne Kontakt im Export, also über den bekannten Weg ansprechen: " +
                     ", ".join("%s (%s)" % (p["firma"], bestand_kurz(p)) if bestand_kurz(p) != "ja" else p["firma"] for p in ohne) + ".")
            z.append("")
    z.append("## So geht es weiter")
    z.append("")
    z.append("- Liste in Excel öffnen, Blatt Kontakte nach Priorität A filtern und Zuordnungen kurz gegenlesen.")
    z.append("- Textbausteine im Blatt Textbausteine anpassen, Platzhalter in eckigen Klammern ersetzen.")
    z.append("- Nachrichten selbst über LinkedIn schicken. Nichts wurde versendet und nichts wird automatisch versendet.")
    z.append("- Bei der ersten Nachricht sagen, woher du den Kontakt hast (wir sind auf LinkedIn vernetzt), "
             "und anbieten, sich nicht wieder zu melden, wenn kein Interesse besteht.")
    z.append("- Datei lokal lassen, nicht in iCloud, OneDrive oder geteilte Ordner legen. Nicht mehr gebrauchte Listen selbst löschen.")
    z.append("")
    with open(pfad, "w", encoding="utf-8") as fobj:
        fobj.write("\n".join(z))


def html_schreiben(pfad, meta, zeilen, segmente, firmen, datum, xlsx_name):
    """Eine Seite zum Anschauen im Browser: Kennzahlen, Verteilung, Top-Firmen, A-Kontakte, bisherige Partner."""
    import html as h

    def e(t):
        return h.escape(str(t if t is not None else ""))

    tab = verteilung(zeilen, segmente)
    bestand = meta.get("bestand") or {}
    mit_bestand = bool(bestand.get("partner"))
    a_kontakte = [z for z in zeilen if z["prioritaet"] == "A"]
    firmen_echt = [f for f in firmen if not f["ohne"]]
    treffer = [f for f in firmen_echt if f["bestand"]]
    im_export = {firma_schluessel(x["firma"]) for x in zeilen}
    ohne = [p for k, p in bestand.get("partner", {}).items() if k not in im_export] if mit_bestand else []

    def badge(prio):
        return '<span class="prio p%s">%s</span>' % (prio, prio)

    def bestand_badge(text):
        return ' <span class="bestand">%s</span>' % e(text) if text else ""

    kacheln = [
        (zahl(len(zeilen)), "Kontakte im Export"),
        (zahl(len(firmen_echt)), "Firmen"),
        (zahl(len(a_kontakte)), "Kontakte mit Priorität A"),
        (zahl(len({firma_schluessel(z["firma"]) for z in a_kontakte if z["firma"]})), "Firmen mit A-Kontakt"),
    ]
    if mit_bestand:
        kacheln.append((zahl(len(treffer)), "bisherige Partner im Export"))

    teile = []
    teile.append("<!doctype html><html lang=\"de\"><head><meta charset=\"utf-8\">")
    teile.append("<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">")
    teile.append("<title>Kontakt-Segmentierung %s</title>" % e(datum_de(datum)))
    teile.append("""<style>
:root{--ink:#1b1b1b;--muted:#6b6b6b;--line:#e3e3e3;--bg:#f7f6f3;--card:#fff;--a:#1f7a4d;--b:#b7791f;--c:#8a8a8a;--accent:#ff4500;--blue:#1f5f8b}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 -apple-system,BlinkMacSystemFont,"Segoe UI",Inter,Arial,sans-serif}
main{max-width:1180px;margin:0 auto;padding:32px 20px 60px}h1{font-size:28px;margin:0 0 4px}h2{font-size:18px;margin:36px 0 12px}
.sub{color:var(--muted);margin:0 0 24px}.kacheln{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}
.kachel{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:16px 18px}.kachel b{display:block;font-size:30px;line-height:1.1}
.kachel span{color:var(--muted);font-size:13px}table{width:100%;border-collapse:collapse;background:var(--card);border:1px solid var(--line);border-radius:12px;overflow:hidden}
th,td{padding:9px 12px;text-align:left;border-bottom:1px solid var(--line);vertical-align:top;font-size:14px}th{background:#f0efeb;font-weight:600;font-size:12.5px;letter-spacing:.02em;text-transform:uppercase;color:var(--muted)}
td.r,th.r{text-align:right}tr:last-child td{border-bottom:none}tr.a td{background:#f2faf5}
.prio{display:inline-block;min-width:24px;text-align:center;border-radius:6px;padding:1px 6px;font-weight:700;font-size:12.5px;color:#fff}
.pA{background:var(--a)}.pB{background:var(--b)}.pC{background:var(--c)}
.bestand{display:inline-block;border-radius:999px;padding:1px 9px;font-size:12px;background:#ddebf7;color:var(--blue);font-weight:600;white-space:nowrap}
.hinweis{background:#fff6ef;border:1px solid #ffd9c2;border-radius:12px;padding:12px 16px;color:#7a3b14;margin-top:28px}
.leer{color:var(--muted);font-style:italic}.balken{display:flex;height:14px;border-radius:7px;overflow:hidden;background:#eee;margin-top:6px}
.balken i{display:block;height:100%}.balken .ia{background:var(--a)}.balken .ib{background:var(--b)}.balken .ic{background:var(--c)}
.legende{color:var(--muted);font-size:12.5px;margin-top:6px}
</style></head><body><main>""")
    teile.append("<h1>Kontakt-Segmentierung, %s</h1>" % e(datum_de(datum)))
    teile.append("<p class=\"sub\">Quelle %s, eingelesen am %s. Vollständige Liste in <code>%s</code>. "
                 "Nichts wurde versendet, jede Nachricht schickst du selbst.</p>"
                 % (e(meta.get("quelle", "?")), e(datum_de(meta.get("eingelesen_am", ""))), e(xlsx_name)))
    teile.append("<div class=\"kacheln\">" + "".join("<div class=\"kachel\"><b>%s</b><span>%s</span></div>" % (e(w), e(t))
                                                   for w, t in kacheln) + "</div>")

    teile.append("<h2>Verteilung nach Segment</h2>")
    teile.append("<table><tr><th>Segment</th><th class=\"r\">A</th><th class=\"r\">B</th><th class=\"r\">C</th><th class=\"r\">Gesamt</th><th>Anteil</th></tr>")
    for seg in segmente:
        t = tab[seg]
        g = sum(t.values())
        if not g:
            continue
        balken = "".join("<i class=\"i%s\" style=\"width:%.1f%%\"></i>" % (p.lower(), 100.0 * t[p] / len(zeilen)) for p in PRIOS)
        teile.append("<tr><td>%s</td><td class=\"r\">%d</td><td class=\"r\">%d</td><td class=\"r\">%d</td><td class=\"r\"><b>%d</b></td>"
                     "<td style=\"width:28%%\"><div class=\"balken\">%s</div></td></tr>" % (e(seg), t["A"], t["B"], t["C"], g, balken))
    summe = {p: sum(tab[seg][p] for seg in segmente) for p in PRIOS}
    teile.append("<tr><td><b>Gesamt</b></td><td class=\"r\"><b>%d</b></td><td class=\"r\"><b>%d</b></td><td class=\"r\"><b>%d</b></td>"
                 "<td class=\"r\"><b>%d</b></td><td></td></tr></table>" % (summe["A"], summe["B"], summe["C"], len(zeilen)))
    teile.append("<p class=\"legende\">Priorität A grün (direkt ansprechen), B gelb (zweite Runde), C grau (nur bei Bedarf).</p>")

    if mit_bestand:
        teile.append("<h2>Bisherige Partner (Reaktivierung)</h2>")
        teile.append("<p class=\"sub\">Partnerliste %s mit %d Firmen. %d davon haben Kontakte im Export.</p>"
                     % (e(bestand.get("quelle", "?")), len(bestand["partner"]), len(treffer)))
        if treffer:
            teile.append("<table><tr><th>Firma</th><th>Bisher</th><th class=\"r\">Kontakte</th><th class=\"r\">davon A</th><th>Wichtigster Kontakt</th><th>Position</th></tr>")
            for f in treffer:
                teile.append("<tr%s><td><b>%s</b></td><td>%s</td><td class=\"r\">%d</td><td class=\"r\">%d</td><td>%s</td><td>%s</td></tr>"
                             % (' class="a"' if f["a"] else "", e(f["firma"]), bestand_badge(f["bestand"]), f["kontakte"], f["a"],
                                e(f["kontakt"]), e(f["position"] or "ohne Position")))
            teile.append("</table>")
        if ohne:
            teile.append("<p class=\"legende\">Ohne Kontakt im Export, also über den bekannten Weg ansprechen: %s.</p>"
                         % e(", ".join("%s (%s)" % (p["firma"], bestand_kurz(p)) if bestand_kurz(p) != "ja" else p["firma"] for p in ohne)))

    teile.append("<h2>Top-20-Firmen</h2>")
    teile.append("<table><tr><th class=\"r\">Rang</th><th>Firma</th><th class=\"r\">Kontakte</th><th class=\"r\">davon A</th><th>Bestes Segment</th><th>Prio</th><th>Wichtigster Kontakt</th><th>Position</th></tr>")
    for i, f in enumerate(firmen_echt[:20], 1):
        teile.append("<tr%s><td class=\"r\">%d</td><td><b>%s</b>%s</td><td class=\"r\">%d</td><td class=\"r\">%d</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>"
                     % (' class="a"' if f["prioritaet"] == "A" else "", i, e(f["firma"]), bestand_badge(f["bestand"]) if mit_bestand else "",
                        f["kontakte"], f["a"], e(f["segment"]), badge(f["prioritaet"]), e(f["kontakt"]), e(f["position"] or "ohne Position")))
    teile.append("</table>")

    teile.append("<h2>Relevante Kontakte (Priorität A, %d)</h2>" % len(a_kontakte))
    if a_kontakte:
        teile.append("<table><tr><th>Name</th><th>Position</th><th>Firma</th><th>Segment</th><th>Begründung</th></tr>")
        for z in a_kontakte:
            teile.append("<tr class=\"a\"><td><b>%s</b></td><td>%s</td><td>%s%s</td><td>%s</td><td>%s</td></tr>"
                         % (e(("%s %s" % (z["vorname"], z["nachname"])).strip()), e(z["position"] or "ohne Position"), e(z["firma"]),
                            bestand_badge(z.get("bestand")) if mit_bestand else "", e(z["segment"]), e(z["begruendung"])))
        teile.append("</table>")
    else:
        teile.append("<p class=\"leer\">Keine Kontakte mit Priorität A.</p>")

    teile.append("<div class=\"hinweis\">Nächster Schritt: Excel öffnen, Blatt Kontakte nach Priorität A filtern, Zuordnung kurz gegenlesen, "
                 "dann Textbausteine anpassen. Nichts wird automatisch versendet. Liste lokal lassen, nicht in Cloud-Ordner legen.</div>")
    teile.append("</main></body></html>")
    with open(pfad, "w", encoding="utf-8") as fobj:
        fobj.write("\n".join(teile))


def reiter_slug(titel):
    s = str(titel or "").strip().lower()
    for alt, neu in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        s = s.replace(alt, neu)
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:64] or "reiter"


def reiter_schreiben(html_pfad, titel, untertitel, quelle, dash_dir=None):
    """Kopiert die HTML-Übersicht als eigenen Reiter ins AIOS-Dashboard. Gibt (slug, zielpfad) zurück."""
    dash = os.path.expanduser(dash_dir or os.path.join("~", ".claude", "dashboard"))
    if not os.path.isdir(os.path.join(dash, "public")):
        fehler("Das AIOS-Dashboard ist nicht eingerichtet (%s fehlt). Einmal /aios-dashboard ausführen, dann erneut." % os.path.join(dash, "public"))
    ordner = os.path.join(dash, "public", "reiter")
    os.makedirs(ordner, exist_ok=True)
    slug = reiter_slug(titel)
    ziel = os.path.join(ordner, slug + ".html")
    shutil.copyfile(html_pfad, ziel)
    with open(os.path.join(ordner, slug + ".json"), "w", encoding="utf-8") as f:
        json.dump({"title": str(titel).strip()[:80], "subtitle": untertitel, "source": quelle,
                   "updated_at": dt.datetime.now().isoformat(timespec="seconds")}, f, ensure_ascii=False, indent=1)
    return slug, ziel


def main():
    ap = argparse.ArgumentParser(description="Segmentierte Kontakte als Excel und Markdown schreiben")
    ap.add_argument("--arbeit", required=True)
    ap.add_argument("--segmente", help="Standard: kontakte/segmente.md, sonst die im Skill-Ordner")
    ap.add_argument("--ausgabe", default=os.path.join("kontakte", "ausgabe"))
    ap.add_argument("--datum", default=dt.date.today().isoformat())
    ap.add_argument("--vorschau", action="store_true")
    ap.add_argument("--cloud-ok", action="store_true",
                    help="Ausgabeordner in iCloud oder OneDrive bewusst zulassen (Nutzer hat entschieden)")
    ap.add_argument("--reiter", metavar="TITEL",
                    help="HTML-Übersicht zusätzlich als Reiter ins AIOS-Dashboard legen, etwa --reiter \"ADAN Karriere\"")
    ap.add_argument("--dashboard", help="Dashboard-Ordner (Standard ~/.claude/dashboard), nur für Tests")
    ap.add_argument("--nur-reiter", action="store_true",
                    help="nichts neu schreiben, nur die neueste HTML-Übersicht aus --ausgabe als Reiter ins Dashboard legen")
    args = ap.parse_args()
    if args.nur_reiter:
        if not args.reiter:
            fehler("--nur-reiter braucht --reiter \"<Titel>\".")
        kandidaten = sorted(glob.glob(os.path.join(args.ausgabe, "segmentierung-*.html")), key=os.path.getmtime)
        if not kandidaten:
            fehler("Keine HTML-Übersicht in %s gefunden. Zuerst Phase A abschließen." % args.ausgabe)
        html = kandidaten[-1]
        slug, ziel = reiter_schreiben(html, args.reiter, "Kontakt-Segmentierung", os.path.abspath(html), args.dashboard)
        print(json.dumps({"reiter": {"titel": args.reiter.strip(), "slug": slug, "datei": ziel, "quelle": html,
                                     "hinweis": "Reiter erscheint im Dashboard in der Gruppe Eigene Reiter, spätestens beim nächsten Datenabruf."}},
                         ensure_ascii=False, indent=1))
        return
    args.segmente = regeldatei("segmente.md", os.path.dirname(os.path.dirname(os.path.abspath(args.arbeit))),
                               args.segmente)

    meta, zeilen, segmente, anrede, texte = laden(args)
    if args.vorschau:
        vorschau(meta, zeilen, segmente, anrede, texte)
        return
    wolke = cloud_hinweis(args.ausgabe)
    if wolke and not args.cloud_ok:
        fehler(wolke + " Die Ergebnisliste würde mitsynchronisiert. Vorschlag: ein lokaler Ordner "
               "wie ~/Kontakte-lokal. Wer trotzdem hier speichern will, startet erneut mit --cloud-ok.", code=3)
    os.makedirs(args.ausgabe, exist_ok=True)
    basis = freier_name(args.ausgabe, "segmentierung-%s" % args.datum)
    firmen = firmen_bilden(zeilen, segmente)
    xlsx = os.path.join(args.ausgabe, basis + ".xlsx")
    md = os.path.join(args.ausgabe, basis + ".md")
    html = os.path.join(args.ausgabe, basis + ".html")
    excel_schreiben(xlsx, meta, zeilen, segmente, anrede, texte, firmen)
    markdown_schreiben(md, meta, zeilen, segmente, firmen, args.datum, basis + ".xlsx")
    html_schreiben(html, meta, zeilen, segmente, firmen, args.datum, basis + ".xlsx")
    bestand = meta.get("bestand") or {}
    aus = {"excel": xlsx, "zusammenfassung": md, "uebersicht_html": html, "kontakte": len(zeilen),
           "firmen": len([f for f in firmen if not f["ohne"]]),
           "bisherige_partner_im_export": len([f for f in firmen if f["bestand"]]) if bestand.get("partner") else None}
    if args.reiter:
        a = sum(1 for z in zeilen if z["prioritaet"] == "A")
        unter = "%s Kontakte, %s mit Priorität A, Quelle %s" % (zahl(len(zeilen)), zahl(a), meta.get("quelle", "Export"))
        slug, ziel = reiter_schreiben(html, args.reiter, unter, os.path.abspath(html), args.dashboard)
        aus["reiter"] = {"titel": args.reiter.strip(), "slug": slug, "datei": ziel,
                         "hinweis": "Reiter erscheint im Dashboard in der Gruppe Eigene Reiter, spätestens beim nächsten Datenabruf."}
    print(json.dumps(aus, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
